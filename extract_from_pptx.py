"""
Step 1 - Pull all photos out of the newsletter PPTX into ./images/
Usage:  python extract_from_pptx.py "Digital Newsletter JAS'26.pptx"
- Every picture is also dumped to ./images/_raw/ (named sXX_YY.png = slide XX, shape YY)
  so you can find a new photo and map it in MAP / TEAM below.
- Existing files in ./images/ are NOT overwritten (your manual photos win).
  Use --force to overwrite.
"""
import io, os, sys
from PIL import Image
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "images"); RAW = os.path.join(IMG, "_raw")

# Team avatars on slide 2 (key = shape id printed in _raw). Matched to the name labels on slide 2.
TEAM = {'minh':'s02_01','dom':'s02_02','xuan':'s02_03','quyen':'s02_04','luan':'s02_05',
        'nghi':'s02_06','thuc':'s02_07','angarish':'s02_22','ky':'s02_30'}
# Content photos -> newsletter slot names
MAP = {'gsmc_recognition':'s17_09','berlin_walkthrough':'s17_12','vac_rtd':'s11_73','td_optimization':'s12_01',
       'rccp_whedds':'s18_06','pebase_timeline':'s09_05','benvoice':'s13_01','benvoice2':'s13_02',
       'mover_apps':'s13_09','mover2':'s13_18','bos_revamp':'s13_19','bos_logo':'s13_10',
       'cyber_foundation':'s06_01','cyber_month':'s07_01','prime_alics':'s18_04',
       'bootcamp_s2_poster':'s15_02','fy2627_emerging_capability':'s15_01',
       'benchmark_vnps':'s14_02','ky_announcement':'s20_04','wise_usecases':'s10_02'}

def walk(shs):
    for sh in shs:
        if sh.shape_type == 6: yield from walk(sh.shapes)
        else: yield sh

def main():
    if len(sys.argv) < 2: print(__doc__); return
    force = "--force" in sys.argv
    os.makedirs(RAW, exist_ok=True)
    prs = Presentation(sys.argv[1]); found = {}
    for si, s in enumerate(prs.slides, 1):
        for j, sh in enumerate(walk(s.shapes)):
            try: blob = sh.image.blob
            except Exception: continue
            im = Image.open(io.BytesIO(blob)); W, H = im.size
            l, t, r, b = sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom
            if any([l, t, r, b]):
                im = im.crop((max(0,int(W*l)), max(0,int(H*t)), min(W,int(W*(1-r))), min(H,int(H*(1-b)))))
            if getattr(sh, "rotation", 0): im = im.rotate(-sh.rotation, expand=True)
            im = im.convert("RGB"); k = f"s{si:02d}_{j:02d}"; found[k] = im
            t = im.copy(); t.thumbnail((1200, 1200)); t.save(os.path.join(RAW, k + ".png"))
    def out(n): return os.path.join(IMG, n + ".jpg")
    for n, k in TEAM.items():
        if k in found and (force or not os.path.exists(out(n))):
            im = found[k]; w, h = im.size; s = int(min(w, h) * 0.78); x = (w - s)//2; y = int(h * 0.02)
            im.crop((x, y, x+s, y+s)).resize((260, 260), Image.LANCZOS).save(out(n), quality=88); print("team ", n)
    for n, k in MAP.items():
        if k in found and (force or not os.path.exists(out(n))):
            im = found[k].copy(); im.thumbnail((1400, 1400), Image.LANCZOS); im.save(out(n), quality=82); print("photo", n)
    print("Done. Raw dump in images/_raw/")

if __name__ == "__main__": main()
